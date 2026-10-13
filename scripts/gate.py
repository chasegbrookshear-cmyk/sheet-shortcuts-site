"""Build gate over data/entries.json (Growth schema v0.1). See SCHEMA.md.
Statuses: PASS | HOLD (cannot be verified / needs a source fix) | MISMATCH (recalc != Growth's expected value) | FLAGGED (single-app, unverifiable by design; shown with a visible flag)
Usage: python3 scripts/gate.py [--all-report]   exit 1 if any entry in data/publish.json is not PASS."""
import json, re, sys, shutil, subprocess, tempfile, datetime
from pathlib import Path
from urllib.parse import urlparse

ROOT = Path(__file__).resolve().parent.parent
OFFICIAL = ("support.microsoft.com", "learn.microsoft.com", "support.google.com")
SECTIONS = {"excel_windows": "sources/ms-win.txt", "excel_mac": "sources/ms-mac.txt",
            "sheets_windows": "sources/g-pc.txt", "sheets_mac": "sources/g-mac.txt"}
SECTION_URL = {"excel": "https://support.microsoft.com/en-us/office/keyboard-shortcuts-in-excel-1798d9d5-842a-42b8-9c99-9b7213f0040f",
               "sheets": "https://support.google.com/docs/answer/181110"}
BANNED = [r"renewal by andersen", r"\bandersen\b", r"ship ?lab", r"window replacement", r"home improvement",
          r"\btax (advice|return|deduction)", r"invest(ing|ment) advice", r"\bHR advice"]
HOLD_LIST = {"regexextract": "Held per Growth/parent: LibreOffice cannot recalc REGEXEXTRACT; needs another check."}
NEEDS_PREFIX = {"XLOOKUP": "_xlfn.XLOOKUP", "TEXTJOIN": "_xlfn.TEXTJOIN", "IFS": "_xlfn.IFS", "LET": "_xlfn.LET",
                "FILTER": "_xlfn._xlws.FILTER", "UNIQUE": "_xlfn.UNIQUE", "SORT": "_xlfn._xlws.SORT",
                "TEXTSPLIT": "_xlfn.TEXTSPLIT", "CONCAT": "_xlfn.CONCAT", "XMATCH": "_xlfn.XMATCH"}
DASH = ("—", "-", "", None)
STOP = set("a an the of to in on and or for with as by is it".split())
# Documented synonym stems (name word -> words Microsoft/Google use for the same action). Keep small; see SCHEMA.md.
SYN = {"link": {"hype"}, "scie": {"expo"}, "last": {"end"}, "top": {"begi", "home"}, "auto": {"sum"}, "jump": {"move"},
       "swit": {"togg"}, "abso": {"abso"}, "acti": {"edit"}}

def load():
    return json.loads((ROOT / "data/entries.json").read_text())
def norm(s): return re.sub(r"\s+", " ", s or "").strip()
def is_dash(v): return v is None or str(v).strip() in DASH

# ---------- shortcut key evidence ----------
WORDS = [("plus sign", "+"), ("minus sign", "-"), ("equal sign", "="), ("hyphen", "-"), ("semicolon", ";"), ("colon", ":"),
         ("grave accent", "`"), ("tilde sign", "~"), ("tilde", "~"), ("dollar sign", "$"), ("percent sign", "%"),
         ("number sign", "#"), ("pound sign", "#"), ("at sign", "@"), ("exclamation point", "!"), ("caret", "^"),
         ("ampersand", "&"), ("underscore", "_"), ("spacebar", "space"), ("command", "cmd"), ("control", "ctrl"),
         ("⌘", "cmd"), ("zero", "0"), ("arrow key", ""), ("arrow", ""), ("key", "")]
def nk(s):
    s = (s or "").lower()
    s = re.sub(r"\b([a-z][a-z ]*?)\s*\(([^\w\s]{1,2})\)", r"\2", s)   # "Semicolon (;)" -> ";"
    s = re.sub(r"\([^)]*\)", "", s)                                     # drop "(with rows selected)"
    s = s.replace("*", "")
    for w, r in WORDS: s = re.sub(rf"(?<![a-z]){re.escape(w)}(?![a-z])", r, s)
    return re.sub(r"\s+", "", s)
def combos(v):
    """'Ctrl+Z / Ctrl+Y' -> two combos; strip Growth annotations like '(or ...)', ', then V', ' opens ...'."""
    out = []
    for part in re.split(r"\s+/\s+", v):
        part = re.split(r"\s+\(|,\s*then\b|\s+opens\b", part)[0].strip()
        if part: out.append(part)
    return out
def stems(s): return {w[:4] for w in re.findall(r"[a-z]+", s.lower()) if w not in STOP and len(w) > 2}
def name_stems(n):
    st = stems(n); return st | {x for w in st for x in SYN.get(w, ())}

_sec_cache = {}
def section_lines(key):
    if key not in _sec_cache:
        _sec_cache[key] = [norm(l).strip(" |") for l in (ROOT / SECTIONS[key]).read_text(errors="ignore").splitlines() if norm(l)]
    return _sec_cache[key]
def find_evidence(sec, combo, name):
    """Return (action_line, key_line, overlap) for the best official line pair whose key line lists this exact combo."""
    target = nk(combo); best = None; lines = section_lines(sec)
    for i, l in enumerate(lines):
        inline = ""
        if ": " in l: l, inline = l.split(": ", 1)
        pieces = [nk(p) for p in re.split(r"\s+or\s+|\s+/\s+|,\s*then\s+", l)]
        if target and target in pieces:
            action = inline or (lines[i - 1] if i else "")
            if i > 1 and (action.startswith("(") or not action): action = lines[i - 2] + " " + action
            ov = len(name_stems(name) & stems(action))
            if best is None or ov > best[2]: best = (action, l, ov)
    return best

def check_shortcut(e, r):
    ev = {}
    for k, sec in SECTIONS.items():
        v = e["keys"].get(k)
        if is_dash(v):
            if not e.get("note"): r["errors"].append(f"{k} is a dash but no note explains why")
            ev[k] = None; continue
        ev[k] = []
        for c in combos(v):
            b = find_evidence(k, c, e["name"])
            if not b: r["hold"].append(f"{k}: '{c}' not found in official section {sec}"); continue
            if b[2] == 0: r["hold"].append(f"{k}: '{c}' found only next to an unrelated action: {b[0]!r}"); continue
            ev[k].append({"combo": c, "quote": f"{b[0]} {b[1]}", "saved_copy": sec, "url": SECTION_URL[k.split('_')[0]]})
    r["evidence"] = ev

# ---------- formula recalc ----------
def grid_cells(meta):
    cells = {}
    for rng, rows in meta["example_grid"].items():
        m = re.match(r"([A-Z]+)(\d+)", rng); c0, r0 = m.group(1), int(m.group(2))
        for ri, row in enumerate(rows):
            for ci, v in enumerate(row):
                if isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v): v = datetime.datetime.fromisoformat(v)
                cells[f"{chr(ord(c0) + ci)}{r0 + ri}"] = v
    return cells
def fns(f): return sorted(set(re.findall(r"([A-Z][A-Z0-9.]*)\(", f or "")))
def xl(f):
    for fn, pre in NEEDS_PREFIX.items(): f = re.sub(rf"(?<![A-Z_.]){fn}\(", pre + "(", f)
    return f
RES = "J2"
SPILL = {}
def is_err(v): return isinstance(v, str) and (v.startswith("#") or v.startswith("Err:"))
def recalc(jobs, cells):
    import openpyxl
    from openpyxl.worksheet.formula import ArrayFormula
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice: raise RuntimeError("LibreOffice (soffice) not on PATH")
    tmp = Path(tempfile.mkdtemp(prefix="sss-recalc-")); wb = openpyxl.Workbook(); wb.remove(wb.active)
    for i, (key, f) in enumerate(jobs):
        ws = wb.create_sheet(f"t{i}")
        for a, v in cells.items(): ws[a] = v
        spill = SPILL.get(key[0])
        if spill == "down": ws[RES] = ArrayFormula("J2:J9", xl(f))
        elif spill == "across": ws[RES] = ArrayFormula("J2:O2", xl(f))
        else: ws[RES] = xl(f)
    (tmp / "in").mkdir(); src = tmp / "in/check.xlsx"; wb.save(src); (tmp / "out").mkdir()
    p = subprocess.run([soffice, f"-env:UserInstallation=file://{tmp}/profile", "--headless", "--norestore", "--nolockcheck",
                        "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(tmp / "out"), str(src)],
                       capture_output=True, text=True, timeout=300)
    res = tmp / "out/check.xlsx"
    if p.returncode or not res.exists(): raise RuntimeError(f"LibreOffice recalc failed: {p.stderr or p.stdout}")
    wb2 = openpyxl.load_workbook(res, data_only=True); out = {}
    for i, (key, f) in enumerate(jobs):
        ws = wb2[f"t{i}"]; down = [ws.cell(row=2 + k, column=10).value for k in range(8)]
        across = [ws.cell(row=2, column=10 + k).value for k in range(6)]
        out[key] = {"value": down[0], "down": [v for v in down if v is not None and not is_err(v)], "across": [v for v in across if v is not None and not is_err(v)]}
    shutil.rmtree(tmp, ignore_errors=True); return out
def fmt(v):
    if isinstance(v, datetime.datetime): return v.date().isoformat()
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v)
def compare(exp, got):
    """Growth writes expected as text. Returns (ok, shown_result)."""
    e = re.sub(r"\s*\((spills|spills across|format as date)\)\s*$", "", str(exp)).strip()
    if " | " in e: vals = got["across"]; parts = [x.strip() for x in e.split("|")]
    elif "; " in e: vals = got["down"]; parts = [x.strip() for x in e.split(";")]
    else: vals = [got["value"]]; parts = [e]
    if "format as date" in str(exp) and isinstance(vals[0], (int, float)):
        vals = [datetime.datetime(1899, 12, 30) + datetime.timedelta(days=vals[0])]
    shown = "; ".join(fmt(v) for v in vals)
    if len(vals) != len(parts): return False, shown
    for p, v in zip(parts, vals):
        if p.endswith("…"):
            if not fmt(v).startswith(p[:-1]): return False, shown
            continue
        try:
            if abs(float(p) - float(v)) > 1e-9: return False, shown
        except (TypeError, ValueError):
            if fmt(v) != p: return False, shown
    return True, shown
def computed(got):
    """Only #NAME?/unknown means LibreOffice lacks the function. #VALUE!, #N/A etc. are real results (=> mismatch)."""
    v = got["value"]
    return not (v is None or (isinstance(v, str) and (v.startswith("#NAME") or "_xlfn" in v or v.startswith("Err:5"))))

def check_formula(e, r, jobs):
    a = e.get("availability")
    for app in ("excel", "sheets"):
        f = e.get(f"{app}_formula")
        if is_dash(f):
            if a == "both": r["errors"].append(f"{app}_formula missing but availability=both")
            continue
        if not f.startswith("="): r["errors"].append(f"{app}_formula must start with =")
        src = [s for s in e["sources"] if (urlparse(s["url"]).hostname or "").endswith("microsoft.com" if app == "excel" else "google.com")]
        if not src: r["errors"].append(f"no official {app} source"); continue
        for fn in fns(f):
            pages = [ROOT / s["cache"] for s in src] + [ROOT / f"sources/functions/{app}-{fn.lower()}.txt"]
            q = [l for pth in pages if pth.is_file() for l in pth.read_text(errors="ignore").splitlines() if fn + "(" in l]
            if not q: r["hold"].append(f"{app}: {fn}( not cited by any saved official {app} page")
            else: r.setdefault("quotes", {})[f"{app}:{fn}"] = norm(q[0])[:200]
        jobs.append(((e["slug"], app), f))

def run(verbose=True):
    meta = load(); entries = meta["entries"]
    pub = json.loads((ROOT / "data/publish.json").read_text())["milestone"]
    report = {}; jobs = []
    for e in entries:
        r = {"type": e["type"], "errors": [], "hold": []}; report[e["slug"]] = r
        for k in ("slug", "type", "category", "name", "seo_title", "meta", "sources", "checked"):
            if not e.get(k): r["errors"].append(f"missing {k}")
        if len(e.get("meta", "")) > 160: r["errors"].append("meta > 160 chars")
        if len(e.get("seo_title", "")) > 60: r["errors"].append("seo_title > 60 chars")
        for s in e.get("sources", []):
            if (urlparse(s.get("url", "")).hostname or "") not in OFFICIAL: r["errors"].append(f"non-official source {s.get('url')}")
            if not (ROOT / s.get("cache", "")).is_file(): r["errors"].append(f"saved copy missing: {s.get('cache')}")
        blob = json.dumps(e)
        for b in BANNED:
            if re.search(b, blob, re.I): r["errors"].append(f"banned content: {b}")
        if e["slug"] in HOLD_LIST: r["hold"].append(HOLD_LIST[e["slug"]])
        if e["type"] == "shortcut": check_shortcut(e, r)
        else: check_formula(e, r, jobs)
    if jobs:
        for e in entries:
            x = str(e.get("expected_result_for_recalc_check", ""))
            if "spills across" in x: SPILL[e["slug"]] = "across"
            elif "spills" in x: SPILL[e["slug"]] = "down"
        cells = grid_cells(meta); got = recalc(jobs, cells)
        byslug = {e["slug"]: e for e in entries}
        for (slug, app), f in jobs:
            e = byslug[slug]; r = report[slug]; g = got[(slug, app)]
            rc = r.setdefault("recalc", {})
            if not computed(g):
                rc[app] = {"computed": False, "raw": fmt(g["value"])}
                if e.get("availability") == "sheets-only" or app == "sheets" and e.get("availability") == "different-names":
                    r.setdefault("flagged", []).append(f"{app}: LibreOffice cannot compute {', '.join(fns(f))}; shown unverified")
                else: r["hold"].append(f"{app}: LibreOffice could not compute ({fmt(g['value'])})")
                continue
            ok, shown = compare(e["expected_result_for_recalc_check"], g)
            rc[app] = {"computed": True, "result": shown, "matches_expected": ok}
            if not ok: r.setdefault("mismatch", []).append(f"{app}: LibreOffice gives {shown!r}, Growth expected {e['expected_result_for_recalc_check']!r}")
    counts = {}
    for slug, r in report.items():
        r["status"] = ("FAIL" if r["errors"] else "MISMATCH" if r.get("mismatch") else "HOLD" if r["hold"]
                       else "FLAGGED" if r.get("flagged") else "PASS")
        counts[r["status"]] = counts.get(r["status"], 0) + 1
        if verbose:
            print(f"{r['status']:8} {slug}")
            for x in r["errors"] + r["hold"] + r.get("mismatch", []) + r.get("flagged", []): print(f"           - {x}")
    bad = [s for s in pub if report.get(s, {}).get("status") != "PASS"]
    stamp = {"checked_with": subprocess.run(["soffice", "--version"], capture_output=True, text=True).stdout.strip(),
             "date": datetime.date.today().isoformat(), "counts": counts, "publish_blocked": bad}
    (ROOT / "data/gate-report.json").write_text(json.dumps({"meta": stamp, "entries": report}, indent=1, ensure_ascii=False, default=str))
    if verbose: print("COUNTS", counts, "\nPUBLISH SET", "OK" if not bad else f"BLOCKED: {bad}")
    return not bad, report, stamp

if __name__ == "__main__":
    ok, _, _ = run(); print("GATE PASS" if ok else "GATE FAIL"); sys.exit(0 if ok else 1)
